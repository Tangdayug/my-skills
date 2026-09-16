#!/usr/bin/env python3
import argparse
import json
import re
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill

ROLE_WORDS = [
    '队长', '队员', '成员', '团队成员', '组长', '组员', '负责人', '姓名', '名单',
    '学生', '同学', '参赛成员', '小组成员', '团队', '成员姓名'
]
ROLE_RE = re.compile(r'(?:' + '|'.join(map(re.escape, sorted(ROLE_WORDS, key=len, reverse=True))) + r')\s*[:：]?')
SEP_RE = re.compile(r'[、,，;；/\\|｜+＋&＆\n\r\t ]+')
BRACKET_RE = re.compile(r'[（(\[【{][^）)\]】}]*[）)\]】}]')
SPACE_RE = re.compile(r'\s+')
CHINESE_NAME_RE = re.compile(r'^[\u3400-\u9fff·]{2,8}$')
ASCII_NAME_RE = re.compile(r'^[A-Za-z][A-Za-z .\-]{1,40}$')


def norm_text(v):
    if v is None:
        return ''
    s = str(v).strip()
    s = s.replace('\u3000', ' ')
    return s


def clean_segment(seg):
    s = norm_text(seg)
    if not s:
        return ''
    s = ROLE_RE.sub('', s)
    s = BRACKET_RE.sub('', s)
    s = s.strip(' ：:、,，;；。.!！?？-—_')
    s = SPACE_RE.sub('', s)
    return s


def classify_candidate(s):
    if not s:
        return 'empty'
    # A compact Chinese personal name is usually 2-4 chars, but keep 5-8 as review rather than discard.
    if CHINESE_NAME_RE.fullmatch(s):
        pure_len = len(s.replace('·', ''))
        if 2 <= pure_len <= 4:
            return 'high_confidence'
        return 'needs_review'
    if ASCII_NAME_RE.fullmatch(s):
        return 'needs_review'
    return 'needs_review'


def extract_from_cell(value):
    raw = norm_text(value)
    if not raw:
        return [], []
    prepared = ROLE_RE.sub('', raw)
    prepared = prepared.replace('：', ':')
    # Colons often remain between labels and values; treat them as separators after labels are removed.
    prepared = prepared.replace(':', '、')
    parts = SEP_RE.split(prepared)
    candidates, review = [], []
    for p in parts:
        p2 = clean_segment(p)
        if not p2:
            continue
        # If a segment contains explanatory text plus a plausible Chinese name, do not guess silently.
        cls = classify_candidate(p2)
        if cls == 'high_confidence':
            candidates.append(p2)
        else:
            review.append(p2)
    return candidates, review


def scan_workbook(path):
    wb = load_workbook(path, data_only=False)
    occurrences = defaultdict(list)
    review_items = []
    nonempty = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                raw = norm_text(cell.value)
                if not raw:
                    continue
                nonempty += 1
                cands, review = extract_from_cell(raw)
                for name in cands:
                    occurrences[name].append({'sheet': ws.title, 'cell': cell.coordinate, 'raw': raw})
                for text in review:
                    review_items.append({'sheet': ws.title, 'cell': cell.coordinate, 'raw': raw, 'candidate': text})
    return occurrences, review_items, nonempty


def write_output(out_path, occurrences, review_items, source_path):
    wb = Workbook()
    ws = wb.active
    ws.title = '整理后名单'
    ws['A1'] = '姓名'
    ws['A1'].font = Font(bold=True)
    names = list(occurrences.keys())
    for i, name in enumerate(names, 2):
        ws.cell(i, 1, name)
    ws.column_dimensions['A'].width = 18

    tr = wb.create_sheet('来源追踪')
    headers = ['姓名', '出现次数', '来源位置', '原始内容']
    for c, h in enumerate(headers, 1):
        tr.cell(1, c, h).font = Font(bold=True)
    for r, name in enumerate(names, 2):
        occs = occurrences[name]
        tr.cell(r, 1, name)
        tr.cell(r, 2, len(occs))
        tr.cell(r, 3, '；'.join(f"{o['sheet']}!{o['cell']}" for o in occs))
        tr.cell(r, 4, ' || '.join(dict.fromkeys(o['raw'] for o in occs)))
    tr.column_dimensions['A'].width = 18
    tr.column_dimensions['B'].width = 10
    tr.column_dimensions['C'].width = 35
    tr.column_dimensions['D'].width = 80

    rv = wb.create_sheet('需要确认')
    for c, h in enumerate(['来源位置', '原始内容', '待确认文本', '说明'], 1):
        rv.cell(1, c, h).font = Font(bold=True)
        rv.cell(1, c).fill = PatternFill('solid', fgColor='FCE4D6')
    if review_items:
        for r, item in enumerate(review_items, 2):
            rv.cell(r, 1, f"{item['sheet']}!{item['cell']}")
            rv.cell(r, 2, item['raw'])
            rv.cell(r, 3, item['candidate'])
            rv.cell(r, 4, '无法高置信判断为标准中文姓名，请人工确认后再纳入最终名单')
    else:
        rv.cell(2, 1, '无')
        rv.cell(2, 4, '未发现自动规则无法判断的内容')
    rv.column_dimensions['A'].width = 22
    rv.column_dimensions['B'].width = 70
    rv.column_dimensions['C'].width = 30
    rv.column_dimensions['D'].width = 55

    audit = wb.create_sheet('核对摘要')
    audit.append(['项目', '结果'])
    audit['A1'].font = audit['B1'].font = Font(bold=True)
    audit.append(['源文件', str(source_path)])
    audit.append(['去重后高置信姓名数', len(names)])
    audit.append(['需要确认项数', len(review_items)])
    duplicate_count = sum(1 for occs in occurrences.values() if len(occs) > 1)
    audit.append(['重复提交姓名数', duplicate_count])
    audit.append(['校验说明', '最终交付前必须人工查看“需要确认”表，并对照“来源追踪”做二次复核。'])
    audit.column_dimensions['A'].width = 24
    audit.column_dimensions['B'].width = 90

    wb.save(out_path)
    return len(names), len(review_items), duplicate_count


def main():
    ap = argparse.ArgumentParser(description='Clean messy Excel name submissions into one-name-per-cell list with traceability.')
    ap.add_argument('input_xlsx')
    ap.add_argument('output_xlsx')
    ap.add_argument('--json-report', default=None)
    args = ap.parse_args()
    src = Path(args.input_xlsx)
    out = Path(args.output_xlsx)
    occ, rev, nonempty = scan_workbook(src)
    n_names, n_review, n_dup = write_output(out, occ, rev, src)
    report = {
        'source': str(src), 'output': str(out), 'nonempty_source_cells': nonempty,
        'unique_high_confidence_names': n_names, 'review_items': n_review,
        'duplicate_names': n_dup,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if args.json_report:
        Path(args.json_report).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')

if __name__ == '__main__':
    main()
